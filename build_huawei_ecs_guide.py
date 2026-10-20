from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/robin/project/carsi-idp/CARSI_IdP_华为云ECS部署实施手册.docx"
NAVY, PALE, GRID = "1F4E79", "F4F7FA", "D9E1F2"

def font(run, size=10.5, bold=False, color="202020", family="Arial Unicode MS"):
    run.font.name = family
    rf = run._element.rPr.rFonts
    rf.set(qn("w:ascii"), family); rf.set(qn("w:hAnsi"), family); rf.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)

def text(p, s, size=10.5, bold=False, color="202020", family="Arial Unicode MS"):
    r = p.add_run(s); font(r, size, bold, color, family); return r

def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), color)

def border(cell):
    tcPr=cell._tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in("w:tcBorders")
    if borders is None: borders=OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=borders.find(qn("w:"+edge))
        if e is None: e=OxmlElement("w:"+edge); borders.append(e)
        e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"6"); e.set(qn("w:color"),GRID)

def margins(cell):
    tcPr=cell._tc.get_or_add_tcPr(); m=tcPr.first_child_found_in("w:tcMar")
    if m is None: m=OxmlElement("w:tcMar"); tcPr.append(m)
    for k in ("top","start","bottom","end"):
        e=m.find(qn("w:"+k))
        if e is None: e=OxmlElement("w:"+k); m.append(e)
        e.set(qn("w:w"),"100"); e.set(qn("w:type"),"dxa")

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); border(c); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,h,9.2,True,"FFFFFF")
        if widths: c.width=Inches(widths[i])
    trPr=t.rows[0]._tr.get_or_add_trPr(); hd=OxmlElement("w:tblHeader"); hd.set(qn("w:val"),"true"); trPr.append(hd)
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; border(c); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n%2: shade(c,PALE)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1
            if i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            text(p,str(val),9.1)
            if widths: c.width=Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def heading(doc, s, level=1):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(15 if level==1 else 10); p.paragraph_format.space_after=Pt(6)
    text(p,s,15 if level==1 else 12.5,True,"000000")
    return p

def body(doc, s, lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.22
    if lead: text(p,lead,10.5,True)
    text(p,s); return p

def bullets(doc, xs):
    for x in xs:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.15; text(p,x)

def code(doc,s):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.2); p.paragraph_format.right_indent=Inches(.12); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(8); p.paragraph_format.line_spacing=1.04
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"F2F2F2"); pPr.append(shd)
    text(p,s,8.2,False,"222222","Courier New")

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles; styles["Normal"].font.name="Arial Unicode MS"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"),"Arial Unicode MS"); styles["Normal"].font.size=Pt(10.5)
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; text(hp,"CARSI IdP 华为云 ECS 部署实施手册",8.5,False,"666666")
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(fp,"部署前请完成机构安全评审与 CARSI 接入审批",8,False,"666666")

p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Inches(.95); p.paragraph_format.space_after=Pt(10); text(p,"CARSI IdP 华为云 ECS 部署实施手册",22,True,"000000")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(24); text(p,"面向正式化准备的 Docker 单节点部署路径",12,False,"4D4D4D")
table(doc,["结论","适用范围"],[
    ("本手册给出可审计的实施顺序", "华为云公有云 ECS  单节点 Docker 部署  以 Ubuntu LTS 为命令示例"),
    ("生产上线前必须补足身份与联邦配置", "域名 DNS  受信 TLS  认证源  属性策略  CARSI 注册  SP 联调"),
    ("云网络必须最小暴露", "公网仅开放 80 和 443  SSH 只允许固定运维源 IP  不开放数据库和 IdP 内部端口"),
], [2.65,4.75])
body(doc,"本手册将现有模拟 IdP 的容器架构迁移为华为云 ECS 上可实施的生产准备方案。文中命令不包含真实口令、私钥、证书或用户目录连接信息；这些值必须通过机构的秘密管理和变更流程提供。", "使用方式：")
body(doc,"本手册假定目标是华为云公有云 Elastic Cloud Server ECS。如使用华为私有云、FusionCompute 或已有校园虚拟化平台，Docker 和 IdP 步骤可复用，但 EIP、VPC、安全组、DNS 控制台和负载均衡步骤应替换为本地平台的等效控制。", "平台边界：")
doc.add_page_break()

heading(doc,"一 目标架构与上线门槛")
body(doc,"推荐先采用单节点 ECS 部署，用于预生产、试运行或规模较小的正式服务。若目标为高可用生产，建议将 Nginx 置于应用弹性负载均衡或 WAF 之后，并将数据库、日志、备份和密钥管理按组织标准拆分。单 ECS 不是高可用架构。")
table(doc,["层次","推荐组件","对外暴露"],[
    ("域名与证书","DNS 公网 Zone  受信 CA 证书  可选 WAF 或 ELB","DNS  A 记录指向 EIP 或负载均衡地址"),
    ("计算","华为云 ECS  Ubuntu LTS  Docker Engine  Compose","仅 TCP 80 与 443"),
    ("应用","Nginx  Shibboleth IdP  MariaDB","容器网络内部 8080 与 3306 不映射宿主机"),
    ("身份与联邦","CAS LDAP AD  CARSI metadata  SP metadata","按最小属性释放与 SP 白名单配置"),
], [1.25,3.45,2.7])
heading(doc,"二 资源与账号准备")
table(doc,["对象","建议或最低要求","实施要点"],[
    ("ECS 规格","起步建议 4 vCPU 8 GiB 内存  100 GiB 系统盘","认证并发和属性解析高时应按压测结果扩容。"),
    ("镜像","Ubuntu 22.04 或 24.04 LTS x86 64","下文命令以 apt 为例；其他发行版须改用对应包管理器。"),
    ("网络","VPC  子网  安全组  EIP","EIP 与 ECS 必须位于同一华为云区域和同一账号。"),
    ("域名","例如 idp.university.edu.cn","先确认机构对 DNS 与证书申请拥有授权。"),
    ("权限","最小 IAM 权限  ECS VPC EIP DNS OBS 或备份服务","普通业务账号不应拥有密钥和网络全局修改权限。"),
], [1.15,2.65,3.6])
body(doc,"华为云安全组同时控制入站与出站流量；放通端口后，操作系统和服务本身仍必须正确监听对应端口。EIP 用于使 ECS 被公网访问，DNS A 记录将域名映射至该公网 IPv4 地址。")

heading(doc,"三 华为云控制台网络配置")
body(doc,"在目标区域创建或选择 VPC 与子网，创建 Linux ECS 并绑定 EIP。不要为 IdP 绑定多个 EIP；若存在多网卡或复杂路由，应由网络管理员评审后再配置。")
table(doc,["安全组方向","协议端口","来源","用途"],[
    ("入方向","TCP 22","固定运维出口 IP 或堡垒机网段","SSH 管理。不得使用 0.0.0.0/0。"),
    ("入方向","TCP 80","0.0.0.0/0 与 IPv6 对应范围","HTTP 到 HTTPS 跳转以及 ACME HTTP 验证。"),
    ("入方向","TCP 443","0.0.0.0/0 与 IPv6 对应范围","SAML 浏览器流量与 IdP metadata 发布。"),
    ("入方向","TCP 3306 8080 8443","不创建公网规则","仅 Docker 内部网络使用。"),
    ("出方向","TCP 443 与必要 DNS NTP","按机构网络策略限制","镜像拉取 证书签发 时间同步及受控 metadata 获取。"),
], [1.1,1.1,2.1,3.1])
body(doc,"在华为云 DNS 创建 Public Zone 后，为 idp 子域新增 A 记录，记录值为 ECS EIP。DNS 缓存传播完成前不要开始生产证书签发或 CARSI metadata 注册。")
code(doc,"# 从任意外部解析器验证  替换为真实域名与 EIP\ndig +short idp.university.edu.cn A\n# 结果必须仅为预期 EIP  再继续证书和应用部署")
doc.add_page_break()

heading(doc,"四 操作系统基线与 Docker 安装")
body(doc,"以 Ubuntu LTS ECS 为例，首先升级系统、创建专用运维用户并确保时间同步。生产环境应优先使用机构的镜像仓库、软件源镜像和基线加固规范；下列命令仅说明必要步骤。")
code(doc,"sudo apt-get update\nsudo apt-get -y upgrade\nsudo apt-get install -y ca-certificates curl gnupg chrony ufw\nsudo systemctl enable --now chrony\ntimedatectl status\nsudo adduser --disabled-password --gecos '' carsi\nsudo usermod -aG sudo carsi")
body(doc,"SSH 应采用密钥认证并在确认新会话可用后关闭密码认证。不要在失去控制台或救援通道前重启 SSH 服务。主机防火墙应与安全组一致，数据库端口和 IdP 内部端口均不应对公网开放。")
code(doc,"sudo ufw default deny incoming\nsudo ufw default allow outgoing\nsudo ufw allow from <运维出口CIDR> to any port 22 proto tcp\nsudo ufw allow 80/tcp\nsudo ufw allow 443/tcp\nsudo ufw enable\nsudo ufw status verbose")
body(doc,"安装 Docker Engine 与 Compose 插件后，使用专用账户执行 Compose。若机构要求使用内部镜像仓库，应先配置 registry mirror 或私有仓库凭据，避免在生产服务器上使用未审查的公共镜像。")
code(doc,"# Ubuntu 软件源示例  生产中可替换为机构批准的 Docker Engine 源\nsudo apt-get install -y docker.io docker-compose-v2\nsudo systemctl enable --now docker\nsudo usermod -aG docker carsi\ndocker --version\ndocker compose version")

doc.add_page_break()
heading(doc,"五 准备部署目录与秘密")
body(doc,"建议在 /srv/carsi-idp 保存部署文件，并将 .env、credentials、证书和备份访问权限限制给 carsi 管理组。不要把 .env、私钥、数据库转储或 credentials 目录提交到 Git。")
code(doc,"sudo install -d -o carsi -g carsi -m 0750 /srv/carsi-idp\nsudo -iu carsi\ncd /srv/carsi-idp\n# 从受控发布包或 CI 制品上传 docker-compose.yml nginx scripts customized mariadb-init\ninstall -d -m 0700 credentials metadata logs mariadb-data nginx/certs nginx/logs\ninstall -m 0600 /dev/null .env")
table(doc,["秘密或文件","生产要求","禁止事项"],[
    (".env","填入真实域名 Scope 管理邮箱和随机口令  chmod 600","使用示例默认密码或发送到聊天群。"),
    ("credentials","首次生成后离线加密备份  仅服务账户可读","升级时重新生成或删除。"),
    ("nginx/certs","受信 CA 证书和私钥  私钥 chmod 600","用临时自签名证书作为正式入口。"),
    ("customized","属性 认证 metadata 覆盖配置经审查后发布","直接覆盖新版 IdP 的全部默认配置。"),
], [1.35,3.2,2.4])
doc.add_page_break()

heading(doc,"六 配置真实域名 证书与 Compose")
body(doc,"将以下值替换为机构实际信息。EntityID、IDP_HOSTNAME、DNS 名称和证书 SAN 必须一致。若 Nginx 与 IdP 前还有 WAF 或 ELB，应确认其将原始 Host 和 X Forwarded Proto 正确传入；否则 SAML 返回地址可能错误。")
code(doc,"IDP_HOSTNAME=idp.university.edu.cn\nIDP_SCOPE=university.edu.cn\nIDP_ENTITY_ID=https://idp.university.edu.cn/idp/shibboleth\nIDP_DISPLAY_NAME_ZH=某大学统一身份认证\nIDP_DISPLAY_NAME_EN=University Identity Provider\nIDP_ADMIN_EMAIL=idp-admin@university.edu.cn\nTZ=Asia/Shanghai\nENABLE_SSL=0\nENABLE_SEALER=0\n# 下列所有口令使用长度不少于 32 的随机值  不写入本文档\nSEALER_STORE_PASSWORD=<secret>\nSEALER_KEY_PASSWORD=<secret>\nBACKCHANNEL_KEYSTORE_PASSWORD=<secret>\nJETTY_KEYSTORE_PASSWORD=<secret>\nPERSISTENT_ID_SALT=<secret>\nMARIADB_ROOT_PASSWORD=<secret>\nMARIADB_PASSWORD=<secret>")
body(doc,"生产 Compose 应绑定 80:80 与 443:443。不要保留模拟环境的 8088:80 与 8443:443 覆盖文件。Nginx server_name 应改为真实域名；在证书、DNS 与安全组未完成前，不要启动占用正式 80 和 443 的服务。")
code(doc,"# 确认证书文件来自机构批准的 CA 或已验证的 ACME 流程\ninstall -m 0644 /secure-transfer/fullchain.pem nginx/certs/idp.crt\ninstall -m 0600 /secure-transfer/privkey.pem nginx/certs/idp.key\n# 修改 nginx/conf.d/idp.conf\n# server_name idp.university.edu.cn;\n# return 301 https://$host$request_uri;\n# Compose 端口使用 80:80 和 443:443")
body(doc,"证书签发可通过机构 PKI、受管证书服务或 ACME 完成。无论使用何种方式，都应在切换前通过 openssl 核验 CN、SAN、颁发者、证书链和有效期。TLS 证书轮换与 SAML 签名证书轮换是两个独立流程。")

doc.add_page_break()
heading(doc,"七 初次启动与技术验收")
code(doc,"cd /srv/carsi-idp\n# 初次且仅初次生成 IdP 签名 加密 sealer 凭据\nbash scripts/gen-credentials.sh\ndocker compose config --quiet\ndocker compose pull\ndocker compose up -d\ndocker compose ps\ndocker inspect carsi-idp --format '{{.State.Health.Status}}'")
body(doc,"只有在首次、credentials 目录为空且已完成安全备份计划时才运行凭据生成脚本。之后的重启、升级、容器重建均应复用同一 credentials 目录。")
code(doc,"curl -fsSI https://idp.university.edu.cn/idp/shibboleth\necho | openssl s_client -connect idp.university.edu.cn:443 \\\n  -servername idp.university.edu.cn 2>/dev/null \\\n  | openssl x509 -noout -subject -issuer -dates -ext subjectAltName\ndocker compose logs --tail=200 idp nginx mariadb")
body(doc,"技术验收通过不等于联邦上线。必须确认 HTTP 跳转、HTTPS、Docker health、IdP 实体标识、证书 SAN 与日志；任何数据库密码、私钥、访问令牌或用户属性不得出现在验收截图与工单正文。")

heading(doc,"八 认证 属性 CARSI 与 SP 联调")
table(doc,["阶段","必须完成的工作","通过条件"],[
    ("认证源","配置并最小权限连接 CAS LDAP AD 或统一身份平台。","测试账户可认证，失败日志不泄露口令或目录详情。"),
    ("属性","定义 attribute resolver 和 filter，按 SP 白名单实施最小释放。","仅释放经批准的属性；无 SP 默认全量属性。"),
    ("CARSI","按 CARSI 流程提交 EntityID metadata 证书 联系人及变更信息。","注册审核通过；发布的 metadata 可被联邦正确解析。"),
    ("SP 联调","导入真实 SP metadata，测试登录 断言 属性 登出和错误页。","至少一个 SP 全链路回归成功，并保留脱敏证据。"),
], [1.05,3.35,2.55])
doc.add_page_break()

heading(doc,"九 备份 监控与升级")
body(doc,"备份至少包含 Compose 文件、.env、customized、credentials、metadata、Nginx 证书和 MariaDB 逻辑转储。备份应加密并复制到 ECS 之外，例如机构备份平台或经过访问控制的 OBS 桶；只把备份留在同一云主机上不能防御主机丢失或误删。")
code(doc,"cd /srv/carsi-idp\ntag=$(date +%Y%m%d-%H%M%S)\nbak=/srv/carsi-idp-backups/${tag}\ninstall -d -m 0700 \"$bak\"\ncp -a .env docker-compose.yml docker-compose.override.yml customized credentials metadata nginx/certs \"$bak/\"\ndocker exec carsi-mariadb sh -c \\\n  'mariadb-dump -uroot -p\"$MARIADB_ROOT_PASSWORD\" --single-transaction --routines --events --all-databases' \\\n  > \"$bak/mariadb-all.sql\"\nchmod 600 \"$bak/mariadb-all.sql\"")
body(doc,"若生产目录没有 docker-compose.override.yml，备份命令应删去该文件参数或先用 test -f 检查。备份任务必须额外验证 SQL 文件非空、签名私钥与 sealer.jks 存在，并定期在隔离环境执行恢复演练。")
table(doc,["变更类型","正确流程","禁止行为"],[
    ("同版本补丁","备份  拉取  先升级 IdP  健康检查  HTTPS 和 SP 回归。","使用 latest 或在未记录镜像 digest 时直接全量重建。"),
    ("IdP 大版本","独立目录与端口演练  合并新默认配置  全量回归  再维护窗口切换。","用旧 customized 全量覆盖新版本或重新生成凭据。"),
    ("TLS 证书","替换 Nginx 证书  nginx -t  重建 Nginx  核验证书链。","误把 TLS 轮换当作 SAML 签名证书轮换。"),
    ("异常回滚","恢复已验证镜像与配置  保留 credentials  从备份恢复数据库。","执行 docker compose down -v 或删除 credentials。"),
], [1.25,3.25,2.45])

heading(doc,"十 上线检查清单")
bullets(doc,[
    "ECS、EIP、VPC、安全组、DNS 和证书所有权均已由机构授权并记录在变更单中。",
    "公网安全组仅允许 SSH 固定源、80、443；3306、8080、8443 不对公网开放。",
    "域名 A 或 AAAA 记录、TLS SAN、IDP_HOSTNAME、EntityID 完全一致。",
    "所有秘密使用随机强值并已写入受控秘密库；.env、证书私钥和 credentials 权限正确。",
    "Docker health、HTTPS、HTTP 跳转、时间同步、日志与数据库健康检查均已通过。",
    "认证源、属性过滤、CARSI metadata、真实 SP 全链路和回滚方案均已完成并由责任人签字。",
    "加密异地备份与恢复演练完成；证书到期、容器异常、磁盘空间和安全更新均有监控。",
])
heading(doc,"十一 华为云操作参考")
body(doc,"华为云控制台界面和区域能力可能调整，实施前应以所选区域、租户和产品版本的官方文档为准。下列链接用于核对 EIP、安全组和 DNS 的控制台操作，而不是替代本手册中的安全评审。")
bullets(doc,[
    "EIP 绑定 ECS  华为云 Elastic IP 文档  https://support.huaweicloud.com/intl/en-us/usermanual-eip/eip_0009.html",
    "安全组规则  华为云 VPC 文档  https://support.huaweicloud.com/intl/en-us/usermanual-vpc/en-us_topic_0030969470.html",
    "公共 DNS 记录集  华为云 DNS 文档  https://support.huaweicloud.com/intl/en-us/usermanual-dns/dns_usermanual_0006.html",
])
body(doc,"文档中的华为云前提基于公开产品文档：EIP 与 ECS 必须处于同一地区和同一账号；安全组规则应只保留所需端口；公共 DNS 的 A 记录用于将域名映射至网站服务器的 IPv4 地址。", "参考说明：")

doc.core_properties.title="CARSI IdP 华为云 ECS 部署实施手册"; doc.core_properties.subject="Docker Shibboleth IdP 部署步骤"; doc.core_properties.author="Codex"
doc.save(OUT); print(OUT)
