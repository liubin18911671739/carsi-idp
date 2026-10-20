from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = "/Users/robin/project/carsi-idp/CARSI_IdP_模拟部署与升级运维说明.docx"

NAVY = "1F4E79"
BLUE = "DCE6F1"
PALE = "F4F7FA"
GRID = "D9E1F2"
TEXT = "202020"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_cell_border(cell, color=GRID):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)

def set_keep(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    pPr.append(keep)

def set_font(run, size=None, bold=None, color=None, name="Arial Unicode MS"):
    run.font.name = name
    rfonts = run._element.rPr.rFonts
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def add_text(p, text, bold=False, size=10.5, color=TEXT):
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return r

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    set_keep(p)
    add_text(p, text, bold=True, size=15 if level == 1 else 12.5, color="000000")
    return p

def add_body(doc, text, lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.22
    if lead:
        add_text(p, lead, bold=True)
    add_text(p, text)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        add_text(p, item, size=10.5)

def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    r = p.add_run(text)
    set_font(r, 8.5, False, "1F1F1F", "Courier New")
    return p

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, label in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, True, 9.4, "FFFFFF")
        if widths: cell.width = Inches(widths[i])
    for n, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if n % 2 == 1: set_cell_shading(cell, PALE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if i == 0 and len(headers) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(p, str(value), size=9.1)
            if widths: cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_step_table(doc, rows):
    return add_table(doc, ["阶段", "操作", "验收或回滚要点"], rows, [1.0, 3.0, 2.4])

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.74)
sec.right_margin = Inches(0.74)

styles = doc.styles
styles["Normal"].font.name = "Arial Unicode MS"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
styles["Normal"].font.size = Pt(10.5)
for name in ("Heading 1", "Heading 2"):
    styles[name].font.name = "Arial Unicode MS"
    styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

# Header/footer
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(header, "CARSI IdP 模拟部署与升级运维说明", size=8.5, color="666666")
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(footer, "内部模拟环境  不可作为 CARSI 正式接入依据", size=8, color="666666")

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Inches(1.05)
p.paragraph_format.space_after = Pt(10)
add_text(p, "CARSI IdP 模拟部署与升级运维说明", True, 23, "000000")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
add_text(p, "远端 Docker 模拟实例  idp.yiyun.edu.cn", size=12, color="4D4D4D")

add_table(doc, ["项目", "当前结论"], [
    ("部署状态", "可用。IdP 与 MariaDB 均通过 Docker 健康检查，HTTPS IdP 端点返回 200。"),
    ("访问隔离", "模拟服务使用 8088 和 8443，未占用宿主机上既有生产服务的 80 和 443。"),
    ("身份标识", "EntityID 为 https://idp.yiyun.edu.cn/idp/shibboleth，scope 为 yiyun.edu.cn。"),
    ("使用边界", "仅验证容器、TLS 反代和 Shibboleth 基础运行；尚未接入真实认证源、CARSI 联邦或真实 SP。"),
], [1.45, 5.95])
add_body(doc, "本说明依据 2026 年 10 月 5 日在 has@100.95.174.122 的实际检查结果编写。它用于后续维护、备份、升级、证书替换和回滚决策；其中的秘密值未写入本文档。", "文档用途：")
add_body(doc, "正式接入前必须完成 DNS、受信 TLS 证书、统一身份认证、属性释放策略、CARSI 注册和真实服务提供方联调。当前模拟实例不得被表述为已完成 CARSI 联邦接入。", "重要边界：")
doc.add_page_break()

add_heading(doc, "一 部署范围与拓扑", 1)
add_body(doc, "部署位于远端目录 /home/has/carsi-idp-sim，使用 Docker Compose 管理。Nginx 负责 TLS 终止和 HTTP 跳转，Shibboleth IdP 提供 /idp/shibboleth 端点，MariaDB 为后续 storage service 与持久化标识扩展提供数据基础。")
add_table(doc, ["请求路径", "服务", "说明"], [
    ("HTTP 8088", "carsi-nginx", "返回 301，并重定向至 https://idp.yiyun.edu.cn:8443。"),
    ("HTTPS 8443", "carsi-nginx", "使用自签名 TLS 证书，反向代理至 IdP 内部 8080。"),
    ("内部 8080", "carsi-idp", "Shibboleth IdP Web 端点，仅在 Docker 网络内暴露。"),
    ("内部 3306", "carsi-mariadb", "MariaDB 数据库，不向宿主机公网发布。"),
], [1.25, 1.45, 4.7])
add_heading(doc, "二 实际运行状态", 1)
add_table(doc, ["服务", "镜像", "端口或网络", "当前状态"], [
    ("carsi-idp", "cscfi/shibboleth-idp:5.2.2", "内部 8080 与 8443", "healthy，restart unless-stopped"),
    ("carsi-nginx", "nginx:1.27-alpine", "0.0.0.0:8088 到 80；0.0.0.0:8443 到 443", "running，restart unless-stopped"),
    ("carsi-mariadb", "mariadb:11.4", "Docker 网络内部 3306", "healthy，restart unless-stopped"),
], [1.1, 2.2, 2.55, 1.55])
add_body(doc, "已完成验证：Docker 健康状态为 healthy；https://idp.yiyun.edu.cn:8443/idp/shibboleth 在指定 Host 解析下返回 200；HTTP 8088 返回 301，并指向隔离的 8443 HTTPS 入口。", "验收结果：")

add_heading(doc, "三 宿主机与容器内系统版本", 1)
add_table(doc, ["范围", "系统或组件", "实测版本或状态"], [
    ("宿主机", "操作系统", "Kylin Linux Advanced Server V11 Swan25"),
    ("宿主机", "内核与架构", "Linux 6.6.0-32.7.v2505.ky11.x86_64，x86_64"),
    ("宿主机", "Docker", "Docker Engine 29.7.2；Docker Compose v5.4.0"),
    ("宿主机", "资源快照", "内存 251 GiB，约 235 GiB 可用；根分区 442 GiB，已用约 108 GiB"),
    ("IdP 容器", "基础系统", "Alpine Linux v3.23.4"),
    ("IdP 容器", "Java 与 Web 容器", "OpenJDK 17.0.18 Alpine；Jetty 12.1.9；Shibboleth IdP 5.2.2"),
    ("Nginx 容器", "基础系统与服务", "Alpine Linux v3.21.3；nginx 1.27.5"),
    ("MariaDB 容器", "基础系统与数据库", "Ubuntu 24.04.5 LTS；MariaDB 11.4.13"),
], [1.15, 2.0, 4.25])
add_body(doc, "镜像标签并不等同于可复现版本。后续生产升级应记录经过验证的镜像 digest，并将其与升级工单、配置快照和验证结果一并归档。")
doc.add_page_break()

add_heading(doc, "四 身份配置与安全资产", 1)
add_table(doc, ["配置项", "当前值或状态", "运维要求"], [
    ("模拟主机名", "idp.yiyun.edu.cn", "当前为模拟名称；正式上线前必须拥有可验证 DNS。"),
    ("EntityID", "https://idp.yiyun.edu.cn/idp/shibboleth", "一旦用于真实联邦，变更将影响 SP 信任关系。"),
    ("scope", "yiyun.edu.cn", "应与正式机构属性命名空间及属性策略一致。"),
    ("TLS 证书", "自签名；CN 与 SAN 均为 idp.yiyun.edu.cn；有效至 2026-11-04", "模拟使用。正式环境必须改为受信 CA 签发证书。"),
    ("运行秘密", ".env 与 credentials 目录中的随机生成口令和密钥", "不得输出、提交版本库或通过普通聊天工具传输。"),
], [1.2, 2.65, 3.55])
add_heading(doc, "五 必须保留的持久化数据", 1)
add_table(doc, ["路径", "内容", "为什么不能随意删除"], [
    ("credentials/", "签名与加密私钥、证书、sealer 密钥库、persistent ID salt", "重新生成会改变信任锚或持久化标识，真实 SP 可能必须重新对接。"),
    ("mariadb-data/", "MariaDB 数据文件", "保存存储服务和后续扩展数据；升级前应有逻辑备份。"),
    ("metadata/", "IdP 发布或导入元数据", "正式接入后是联邦配置的重要状态。"),
    ("customized/conf/", "对 IdP 默认配置的覆盖", "当前包含 EntityID 与 scope；大版本升级应重新合并。"),
    ("nginx/certs/", "反向代理证书和私钥", "TLS 轮换资产，与 SAML 签名证书不同。"),
    (".env", "数据库、sealer 与 Jetty 口令等秘密", "必须保留安全副本，权限应为 600。"),
], [1.55, 2.25, 3.6])
add_body(doc, "严禁执行 docker compose down -v，也不要为排障而删除 credentials 或 mariadb-data。先检查日志、配置、口令匹配和权限；删除持久化资产并不是有效的修复方式。", "禁止操作：")

add_heading(doc, "六 当前模拟环境的未完成事项", 1)
add_bullets(doc, [
    "未配置 idp.yiyun.edu.cn 的真实 DNS 解析和受信任的公网证书。",
    "未接入 CAS、LDAP、Active Directory 或其他真实认证源；因此没有实际用户认证链路。",
    "未导入 CARSI 联邦元数据，未完成 CARSI 注册、审核或生产 metadata 发布。",
    "未配置属性解析、属性过滤、最小化属性释放和真实 SP 属性映射。",
    "未与任何真实 SP 完成 SAML 登录、断言、单点登出和错误恢复联调。",
    "未建立异地加密备份、日志集中审计、证书过期告警和定期恢复演练。",
])
doc.add_page_break()

add_heading(doc, "七 升级前的标准操作", 1)
add_body(doc, "无论是镜像补丁、证书轮换还是 IdP 大版本升级，都必须先备份、再验证、后切换。以下命令仅在当前模拟目录中执行，且不会显示秘密内容。")
add_code(doc, "ssh has@100.95.174.122\ncd /home/has/carsi-idp-sim\ndocker compose ps\ndocker inspect carsi-idp --format '{{.State.Health.Status}}'\ncurl -k --resolve idp.yiyun.edu.cn:8443:127.0.0.1 \\\n  https://idp.yiyun.edu.cn:8443/idp/shibboleth")
add_code(doc, "backup_tag=$(date +%Y%m%d-%H%M%S)\nbackup_dir=\"/home/has/carsi-idp-backups/${backup_tag}\"\nmkdir -p \"${backup_dir}\" && chmod 700 \"${backup_dir}\"\ncp -a .env docker-compose.yml docker-compose.override.yml customized \\\n  credentials metadata nginx/certs \"${backup_dir}/\"\ndocker exec carsi-mariadb sh -c \\\n  'mariadb-dump -uroot -p\"$MARIADB_ROOT_PASSWORD\" --single-transaction --routines --events --all-databases' \\\n  > \"${backup_dir}/mariadb-all.sql\"\nchmod 600 \"${backup_dir}/mariadb-all.sql\"")
add_body(doc, "备份后必须在另一台受控存储中保存副本，并检查 SQL 文件非空、idp-signing.key、idp-encryption.key 与 sealer.jks 均存在。仅生成备份命令或仅看到目录创建成功都不等同于备份可恢复。")

add_heading(doc, "八 同一大版本内的补丁升级", 1)
add_step_table(doc, [
    ("1 评估", "阅读目标镜像版本说明，记录当前 image digest，确认 Java、Jetty 与 MariaDB 兼容性。", "不使用浮动 latest 标签；保留可回退的当前标签或 digest。"),
    ("2 备份", "执行第七节的配置、凭据、证书和数据库备份。", "确认秘密资产和数据库逻辑备份均真实存在。"),
    ("3 拉取", "执行 docker compose pull，并用 docker compose config 检查最终生效配置。", "不在未审查 Compose 合并结果时直接重建全部服务。"),
    ("4 分批", "先 docker compose up -d --no-deps idp，观察日志和健康检查。", "IdP 不健康时停止后续组件升级，回退镜像标签。"),
    ("5 验收", "检查健康状态、HTTPS 200、Nginx 跳转、容器日志与真实联调用例。", "生产环境必须增加至少一个 SP 的登录和属性断言回归。"),
])
add_code(doc, "docker compose pull\ndocker compose config\ndocker compose up -d --no-deps idp\ndocker compose logs --tail=200 idp\ndocker inspect carsi-idp --format '{{.State.Health.Status}}'")

add_heading(doc, "九 Shibboleth IdP 大版本升级", 1)
add_body(doc, "从 IdP 5 升级到未来大版本时，不能把镜像标签直接改掉后上线。应创建独立的升级副本，使用不同 Compose project 名称、容器名和端口，例如 9443 与 9088，并在副本中恢复安全副本进行演练。")
add_bullets(doc, [
    "先阅读目标 IdP 版本的官方升级、Java、Jetty、Servlet API、模块弃用和数据库兼容说明。",
    "从目标镜像导出新版默认 idp.properties，并把现有 EntityID、scope 和必要自定义项逐项合并；不要用旧整份文件覆盖新默认配置。",
    "复制并保留原 credentials、metadata 和数据库备份；绝不运行凭据生成脚本。",
    "比较升级前后 metadata 的 EntityID、签名证书和加密证书。未执行正式轮换时，证书公钥不得意外改变。",
    "完成认证、属性释放、SAML 断言、登出、SP 联调、故障重启和回滚演练后，再在维护窗口切换入口。",
])
doc.add_page_break()

add_heading(doc, "十 TLS 证书轮换", 1)
add_body(doc, "Nginx TLS 证书和 Shibboleth 的 SAML 签名证书是两套不同资产。替换 nginx/certs/idp.crt 与 nginx/certs/idp.key 只影响 HTTPS 信任，不应改变 IdP 的 SAML 签名与加密密钥。")
add_code(doc, "cd /home/has/carsi-idp-sim\ndocker compose exec nginx nginx -t\ndocker compose up -d --force-recreate nginx\necho | openssl s_client -connect 127.0.0.1:8443 \\\n  -servername idp.yiyun.edu.cn 2>/dev/null \\\n  | openssl x509 -noout -subject -issuer -dates -ext subjectAltName")
add_body(doc, "SAML 签名证书轮换需要更严格的双证书发布、SP metadata 更新、宽限期和回收流程，不能用 TLS 证书轮换步骤替代。")

add_heading(doc, "十一 故障回滚原则", 1)
add_step_table(doc, [
    ("镜像升级失败", "将 Compose 中镜像恢复为已验证标签或 digest，执行 docker compose up -d。", "检查 IdP health、HTTPS 端点和日志；不要删除数据卷。"),
    ("配置变更失败", "从升级前备份恢复 .env、customized、Compose 文件和 Nginx 证书。", "恢复后先 docker compose config，再重建受影响服务。"),
    ("数据库迁移失败", "停止写入路径后，从经验证的 SQL 备份或受控数据目录恢复。", "应先在隔离副本验证恢复，不要直接覆盖唯一生产数据。"),
    ("密钥或证书异常", "优先还原原 credentials；核对口令、文件权限和版本兼容性。", "不可通过删除 credentials 解决问题。真实联邦中密钥变更会影响 SP 信任。"),
])

add_heading(doc, "十二 正式化前的建议", 1)
add_bullets(doc, [
    "将远端实际覆盖项回写到受版本控制的部署仓库，消除本地骨架与远端运行态之间的漂移。",
    "为真实域名配置 DNS、受信 TLS 证书、80 和 443 的独立虚拟主机，并完成外网和校内网连通性验证。",
    "实施最小属性释放策略，建立属性清单、SP 白名单、变更审批和审计记录。",
    "建立加密离线备份、证书到期告警、镜像漏洞管理和至少一次完整恢复演练。",
    "在 CARSI 正式注册和真实 SP 联调完成之前，所有对外材料都应将该系统标记为模拟或预生产环境。",
])

doc.core_properties.title = "CARSI IdP 模拟部署与升级运维说明"
doc.core_properties.subject = "远端 Docker 模拟实例部署总结和升级操作规程"
doc.core_properties.author = "Codex"
doc.save(OUT)
print(OUT)
